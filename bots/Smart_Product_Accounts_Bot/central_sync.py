from __future__ import annotations
import json, os, time
from datetime import datetime
from typing import Any
import gspread
from google.oauth2.service_account import Credentials
from openpyxl.utils import get_column_letter
from central_offline import flush_snapshot, sync_snapshot

SCOPES=["https://www.googleapis.com/auth/spreadsheets"]

def enabled() -> bool:
    return bool(os.getenv("CENTRAL_SHEET_ID", "").strip() and (os.getenv("CENTRAL_GOOGLE_CREDENTIALS_JSON", "").strip() or os.getenv("GOOGLE_CREDENTIALS_JSON", "").strip()))

def _book():
    sid=os.getenv("CENTRAL_SHEET_ID", "").strip()
    raw=os.getenv("CENTRAL_GOOGLE_CREDENTIALS_JSON", "").strip() or os.getenv("GOOGLE_CREDENTIALS_JSON", "").strip()
    if not sid or not raw: raise RuntimeError("CENTRAL_SHEET_ID / credentials missing")
    info=json.loads(raw)
    creds=Credentials.from_service_account_info(info, scopes=SCOPES)
    return gspread.authorize(creds).open_by_key(sid)

def _retry(fn,*args,attempts=3,**kwargs):
    last=None
    for i in range(attempts):
        try: return fn(*args,**kwargs)
        except Exception as e:
            last=e
            if i+1<attempts: time.sleep(i+1)
    raise last

def _ws(title:str, headers:list[str]):
    book=_retry(_book)
    try: ws=book.worksheet(title)
    except gspread.WorksheetNotFound: ws=book.add_worksheet(title=title, rows=2000, cols=max(10,len(headers)))
    if ws.row_values(1)[:len(headers)] != headers:
        ws.clear(); ws.update(range_name=f"A1:{get_column_letter(len(headers))}1", values=[headers])
    try: ws.freeze(rows=1)
    except Exception: pass
    return ws

def _replace_tab_online(title:str, headers:list[str], rows:list[list[Any]]) -> int:
    if not enabled(): return 0
    ws=_ws(title,headers)
    ws.clear(); ws.update(range_name=f"A1:{get_column_letter(len(headers))}1", values=[headers])
    if rows: ws.append_rows(rows, value_input_option="USER_ENTERED")
    _touch(title, len(rows))
    return len(rows)

def _touch(source:str,count:int):
    headers=["Source","Last Sync","Rows"]
    ws=_ws("Sync_Status",headers)
    values=ws.get_all_values()
    now=datetime.now().isoformat(timespec="seconds")
    found=None
    for idx,row in enumerate(values[1:],start=2):
        if row and row[0]==source: found=idx; break
    if found: ws.update(range_name=f"A{found}:C{found}", values=[[source,now,count]])
    else: ws.append_row([source,now,count], value_input_option="USER_ENTERED")


def flush_pending_central() -> dict[str, Any]:
    return flush_snapshot("Courier_Monthly", _replace_tab_online)

def read_tab(title:str) -> list[dict[str,Any]]:
    if not enabled(): return []
    try: return _book().worksheet(title).get_all_records()
    except gspread.WorksheetNotFound: return []

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from settlement import available_settlement_months, settlement_summaries

COURIER_HEADERS=["Month","Total Parcel","Total COD","Delivered","Delivered COD","Returned","Return COD","Pending","Pending COD","Courier Charge","Expected Settlement","Paid Invoice Payout","Collected","Remaining","Updated At"]

def sync_courier_snapshot() -> dict[str,Any]:
    months=[(y,m) for y,m,_ in available_settlement_months(120)]
    summaries=settlement_summaries(months)
    now=datetime.now().isoformat(timespec='seconds')
    rows=[]
    for y,m in months:
        s=summaries.get((y,m))
        if not s: continue
        rows.append([f"{y:04d}-{m:02d}",s['total_parcels'],s['total_cod'],s['delivered_parcels'],s['delivered_cod'],s['return_parcels'],s['return_cod'],s['pending_parcels'],s['pending_cod'],s['total_charge'],s['expected_settlement'],s['matched_payout'],s['matched_collected'],s['remaining'],now])
    result = sync_snapshot(
        "Courier_Monthly", "Courier_Monthly", COURIER_HEADERS, rows, _replace_tab_online
    )
    return {
        "synced": int(result.get("synced") or 0),
        "months": len(rows),
        "queued": bool(result.get("queued")),
        "pending": int(result.get("pending") or 0),
        "sync_error": result.get("error") or "",
    }

def export_central_xlsx(path:Path) -> Path:
    courier=read_tab("Courier_Monthly"); bank=read_tab("Dollar_Payments"); expenses=read_tab("Business_Expenses"); status=read_tab("Sync_Status")
    wb=Workbook(); wb.remove(wb.active)
    def add(name, data):
        ws=wb.create_sheet(name)
        if not data:
            ws.append(["No data synced"]); return
        headers=list(data[0].keys()); ws.append(headers)
        for c in ws[1]: c.font=Font(bold=True)
        for row in data: ws.append([row.get(h,"") for h in headers])
        for col in ws.columns:
            width=min(35,max(12,max(len(str(c.value or '')) for c in col)+2)); ws.column_dimensions[col[0].column_letter].width=width
    add("Courier",courier); add("Dollar Payments",bank); add("Business Expenses",expenses); add("Sync Status",status)
    # Summary
    ws=wb.create_sheet("Combined Summary",0); ws.append(["BURAQ CENTRAL ACCOUNTS"]); ws['A1'].font=Font(bold=True,size=16)
    total_expected=sum(float(r.get('Expected Settlement') or 0) for r in courier)
    total_payout=sum(float(r.get('Paid Invoice Payout') or 0) for r in courier)
    dollar=sum(float(r.get('Amount') or 0) for r in bank)
    expense=sum(float(r.get('Amount') or 0) for r in expenses)
    ws.append(["Expected Courier Settlement",total_expected]); ws.append(["Courier Payout",total_payout]); ws.append(["Dollar Payment",dollar]); ws.append(["Business Expense",expense]); ws.append(["Balance after expenses",total_payout-dollar-expense])
    ws.column_dimensions['A'].width=30; ws.column_dimensions['B'].width=20
    wb.save(path); return path
