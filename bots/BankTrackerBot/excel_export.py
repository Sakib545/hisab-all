from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import EXPORTS_DIR

NAVY = "14213D"
BLUE = "1F4E78"
WHITE = "FFFFFF"
LIGHT = "EAF2F8"
GREEN = "D9EAD3"
DARK = "111827"
GRID = Side(style="thin", color="4B5563")
MEDIUM = Side(style="medium", color="111827")
THIN_BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)


def _amount(value):
    return float(value or 0)


def make_excel(rows, payment_code=None):
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Payment Report"
    ws.sheet_view.showGridLines = False

    include_payment = payment_code is None
    headers = (["Payment ID"] if include_payment else []) + [
        "SL", "Date", "Time", "Receiver Account", "Amount"
    ]
    last_col = len(headers)
    total = sum(_amount(r["amount"]) for r in rows)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws.cell(1, 1, "BURAQ DOLLER PAYMENT")
    title.font = Font(size=19, bold=True, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(horizontal="center", vertical="center")
    title.border = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)
    ws.row_dimensions[1].height = 38

    summary = [
        ("Payment ID", payment_code or "ALL PAYMENTS"),
        ("Transactions", len(rows)),
        ("Grand Total", total),
        ("Generated", datetime.now().strftime("%d-%b-%Y %I:%M %p")),
    ]
    for row_no, (label, value) in enumerate(summary, start=2):
        ws.cell(row_no, 1, label)
        ws.cell(row_no, 2, value)
        ws.cell(row_no, 1).font = Font(bold=True, color=DARK)
        ws.cell(row_no, 1).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(row_no, 2).fill = PatternFill("solid", fgColor="FFFFFF")
        for col in range(1, last_col + 1):
            ws.cell(row_no, col).border = THIN_BORDER
        if row_no == 4:
            ws.cell(row_no, 2).number_format = '৳#,##0.00'
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_col)
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=last_col)
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=last_col)
    ws.merge_cells(start_row=5, start_column=2, end_row=5, end_column=last_col)

    header_row = 7
    for col, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[header_row].height = 27

    for sl, r in enumerate(rows, start=1):
        row_no = header_row + sl
        values = ([r["payment_code"] or ""] if include_payment else []) + [
            sl, r["transaction_date"] or "NOT DETECTED", r["transaction_time"] or "NOT DETECTED",
            str(r["account_number"] or "NOT DETECTED"), _amount(r["amount"]),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_no, col, value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if sl % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(row_no, last_col).number_format = '৳#,##0.00'
        ws.cell(row_no, last_col).alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[row_no].height = 23

    data_last = header_row + len(rows)
    if rows:
        ref = f"A{header_row}:{get_column_letter(last_col)}{data_last}"
        table = Table(displayName="BuraqPaymentTable", ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=False, showColumnStripes=False,
        )
        ws.add_table(table)

    total_row = data_last + 2
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=last_col - 1)
    ws.cell(total_row, 1, "GRAND TOTAL")
    ws.cell(total_row, last_col, total)
    for col in range(1, last_col + 1):
        cell = ws.cell(total_row, col)
        cell.font = Font(bold=True, color=DARK)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.border = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)
    ws.cell(total_row, 1).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(total_row, last_col).number_format = '৳#,##0.00'
    ws.cell(total_row, last_col).alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[total_row].height = 26

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{data_last}"
    widths = [15, 8, 15, 18, 25, 18] if include_payment else [8, 15, 18, 25, 18]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"1:{header_row}"
    ws.oddFooter.center.text = "BURAQ DOLLER PAYMENT"
    ws.oddFooter.right.text = "Page &P of &N"

    suffix = payment_code.replace("-", "_") if payment_code else "all_payments"
    path = EXPORTS_DIR / datetime.now().strftime(f"BURAQ_{suffix}_%Y%m%d_%H%M%S.xlsx")
    wb.save(path)
    return path
